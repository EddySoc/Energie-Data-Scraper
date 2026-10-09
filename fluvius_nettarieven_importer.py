"""
Fluvius net tariff importer

Downloads the latest Fluvius distribution network tariff files (electricity and gas)
per Fluvius zone and exports a normalized CSV consumed by the dashboard builder.
"""

from __future__ import annotations

import argparse
from datetime import datetime
import io
import re
from pathlib import Path

import pandas as pd
import requests
from openpyxl import load_workbook

ELEC_PAGE = "https://www.fluvius.be/nl/factuur-en-tarieven/netkosten/netkosten-afname-elektriciteit"
GAS_PAGE = "https://www.fluvius.be/nl/factuur-en-tarieven/netkosten/netkosten-afname-aardgas"
OUTPUT_CSV = "fluvius_nettarieven_latest.csv"
DEFAULT_DATA_DIR = Path(__file__).resolve().parent / "scraper_data"

ZONE_SLUG_TO_NAME = {
    "fluvius-antwerpen": "Fluvius Antwerpen",
    "fluvius-halle-vilvoorde": "Fluvius Halle-Vilvoorde",
    "fluvius-imewo": "Fluvius Imewo",
    "fluvius-kempen": "Fluvius Kempen",
    "fluvius-limburg": "Fluvius Limburg",
    "fluvius-midden-vlaanderen": "Fluvius Midden-Vlaanderen",
    "fluvius-west": "Fluvius West",
    "fluvius-zenne-dijle": "Fluvius Zenne-Dijle",
}

HEADERS = {"User-Agent": "Mozilla/5.0"}


def _to_float(value, *, required: bool = False) -> float:
    if value is None or (isinstance(value, str) and not value.strip()):
        if required:
            raise ValueError("Missing required rate in Fluvius tariff workbook")
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        if required:
            raise ValueError(f"Invalid rate in Fluvius tariff workbook: {value!r}") from exc
        return 0.0


def _find_publication_links(page_url: str, fuel: str) -> dict[str, tuple[str, str]]:
    """Return latest publication URL per zone slug for the requested fuel type."""
    html = requests.get(page_url, headers=HEADERS, timeout=45).text
    fuel_token = "elektriciteit" if fuel == "elec" else "aardgas"

    pattern = re.compile(
        rf'href="(?P<href>/nl/publicatie/(?P<slug>fluvius-[a-z\-]+)-distributienettarief-afname-{fuel_token}-(?P<start>\d{{8}})-(?P<end>\d{{8}}))"',
        flags=re.IGNORECASE,
    )

    latest: dict[str, tuple[str, str]] = {}
    for match in pattern.finditer(html):
        slug = match.group("slug").lower()
        if slug not in ZONE_SLUG_TO_NAME:
            continue

        end_date = match.group("end")
        href = match.group("href")
        pub_url = f"https://www.fluvius.be{href}"

        current = latest.get(slug)
        if current is None or end_date > current[0]:
            latest[slug] = (end_date, pub_url)

    return latest


def _extract_xlsx_url(publication_url: str) -> str:
    html = requests.get(publication_url, headers=HEADERS, timeout=45).text
    m = re.search(
        r'href="(?P<href>(?:https?://[^\"]+|/[^\"]+)\.xlsx(?:\?[^\"]*)?)"',
        html,
        flags=re.IGNORECASE,
    )
    if not m:
        raise ValueError(f"No XLSX link found on publication page: {publication_url}")
    href = m.group("href")
    if href.startswith("http"):
        return href
    return f"https://www.fluvius.be{href}"


def _download_xlsx(url: str) -> bytes:
    response = requests.get(url, headers=HEADERS, timeout=45)
    response.raise_for_status()
    return response.content


def _parse_electricity_rates(xlsx_data: bytes) -> dict[str, float]:
    ws = load_workbook(io.BytesIO(xlsx_data), data_only=True).active
    if ws is None:
        raise ValueError("Electricity workbook has no active worksheet")

    # Low-voltage, peak-meter column. The distribution kWh charge and the
    # separately listed public-service obligations/surcharges are distinct.
    distribution = _to_float(ws.cell(row=17, column=14).value, required=True)
    capacity = _to_float(ws.cell(row=15, column=14).value)
    meter = _to_float(ws.cell(row=27, column=14).value)
    odv_day = _to_float(ws.cell(row=30, column=14).value)
    odv_night = _to_float(ws.cell(row=31, column=14).value)
    surcharges = _to_float(ws.cell(row=33, column=14).value)
    minimum_text = " ".join(
        str(cell.value)
        for row in ws.iter_rows()
        for cell in row
        if cell.value is not None
    )
    minimum_match = re.search(
        r"minimale bijdrage van\s+([\d.,]+)\s+euro/jaar",
        minimum_text,
        flags=re.IGNORECASE,
    )
    if not minimum_match:
        raise ValueError("Could not find the low-voltage minimum annual contribution in the electricity workbook")
    minimum_charge = float(minimum_match.group(1).replace(",", "."))

    return {
        "dist_day_eur_kwh": distribution,
        "dist_night_eur_kwh": distribution,
        "capacity_eur_kw_year": capacity,
        "meter_eur_year": meter,
        "databeheer_eur_year": meter,
        "heffingen_odv_eur_kwh": odv_day,
        "heffingen_odv_night_eur_kwh": odv_night,
        "heffingen_toeslagen_eur_kwh": surcharges,
        "minimum_network_eur_year": minimum_charge,
    }


def _parse_gas_rates(xlsx_data: bytes) -> dict[str, float]:
    ws = load_workbook(io.BytesIO(xlsx_data), data_only=True).active
    if ws is None:
        raise ValueError("Gas workbook has no active worksheet")

    # The four non-telemetered household categories are selected by annual
    # consumption in the dashboard (columns D:G in the tariff workbook).
    tiers = {}
    for tier, column in enumerate(range(4, 8), start=1):
        tiers[f"t{tier}"] = {
            "dist_eur_kwh": _to_float(ws.cell(row=14, column=column).value, required=True),
            "fixed_term_eur_year": _to_float(ws.cell(row=13, column=column).value, required=True),
            "odv_eur_kwh": _to_float(ws.cell(row=24, column=column).value),
            "other_levies_eur_kwh": (
                _to_float(ws.cell(row=27, column=column).value)
                + _to_float(ws.cell(row=28, column=column).value)
            ),
        }
    data_meter = _to_float(ws.cell(row=22, column=4).value)
    result = {
        "gas_databeheer_eur_year": data_meter,
        "gas_meter_eur_year": data_meter,
    }
    for tier, rates in tiers.items():
        for name, value in rates.items():
            result[f"gas_{name}_{tier}"] = value
    # Keep the old flat fields for consumers that have not adopted tiered rates.
    result.update({
        "gas_dist_eur_kwh": tiers["t1"]["dist_eur_kwh"],
        "gas_fixed_term_eur_year": tiers["t1"]["fixed_term_eur_year"],
        "gas_heffingen_odv_eur_kwh": tiers["t1"]["odv_eur_kwh"],
        "gas_heffingen_overig_eur_kwh": tiers["t1"]["other_levies_eur_kwh"],
    })
    return result


def update_fluvius_rates(output_csv: str | Path = DEFAULT_DATA_DIR / OUTPUT_CSV) -> str:
    output_path = Path(output_csv)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    elec_publications = _find_publication_links(ELEC_PAGE, fuel="elec")
    gas_publications = _find_publication_links(GAS_PAGE, fuel="gas")

    if not elec_publications or not gas_publications:
        raise RuntimeError("Could not find Fluvius publication links for electricity and gas.")

    rows: list[dict[str, object]] = []
    for slug, zone_name in ZONE_SLUG_TO_NAME.items():
        if slug not in elec_publications or slug not in gas_publications:
            print(f"[WARN] Missing publication link for zone: {zone_name}")
            continue

        elec_pub_url = elec_publications[slug][1]
        gas_pub_url = gas_publications[slug][1]

        elec_xlsx_url = _extract_xlsx_url(elec_pub_url)
        gas_xlsx_url = _extract_xlsx_url(gas_pub_url)

        elec_rates = _parse_electricity_rates(_download_xlsx(elec_xlsx_url))
        gas_rates = _parse_gas_rates(_download_xlsx(gas_xlsx_url))

        rows.append(
            {
                "zone": zone_name,
                "elec_dist_day_eur_kwh": elec_rates["dist_day_eur_kwh"],
                "elec_dist_night_eur_kwh": elec_rates["dist_night_eur_kwh"],
                "elec_capacity_eur_kw_year": elec_rates["capacity_eur_kw_year"],
                "elec_meter_eur_year": elec_rates["meter_eur_year"],
                "elec_databeheer_eur_year": elec_rates["databeheer_eur_year"],
                "elec_heffingen_odv_eur_kwh": elec_rates["heffingen_odv_eur_kwh"],
                "elec_heffingen_odv_night_eur_kwh": elec_rates["heffingen_odv_night_eur_kwh"],
                "elec_heffingen_toeslagen_eur_kwh": elec_rates["heffingen_toeslagen_eur_kwh"],
                "elec_min_network_eur_year": elec_rates["minimum_network_eur_year"],
                "gas_dist_eur_kwh": gas_rates["gas_dist_eur_kwh"],
                "gas_fixed_term_eur_year": gas_rates["gas_fixed_term_eur_year"],
                "gas_databeheer_eur_year": gas_rates["gas_databeheer_eur_year"],
                "gas_meter_eur_year": gas_rates["gas_meter_eur_year"],
                "gas_heffingen_odv_eur_kwh": gas_rates["gas_heffingen_odv_eur_kwh"],
                "gas_heffingen_overig_eur_kwh": gas_rates["gas_heffingen_overig_eur_kwh"],
                **{
                    key: value
                    for key, value in gas_rates.items()
                    if key.endswith(("_t1", "_t2", "_t3", "_t4"))
                },
                "elec_source": elec_xlsx_url,
                "gas_source": gas_xlsx_url,
                "imported_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            }
        )

    if not rows:
        raise RuntimeError("No Fluvius zones were imported.")

    df = pd.DataFrame(rows)
    df = df.sort_values("zone").reset_index(drop=True)
    df.to_csv(output_path, index=False)
    return str(output_path)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Download and normalize Fluvius network tariffs.")
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=DEFAULT_DATA_DIR,
        help="Directory for the generated Fluvius tariff CSV.",
    )
    args = parser.parse_args()
    out = update_fluvius_rates(args.output_dir / OUTPUT_CSV)
    print(f"[OK] Fluvius nettarieven bijgewerkt: {Path(out).resolve()}")
